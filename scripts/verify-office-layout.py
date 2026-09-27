"""Independently verify actual exports from office-layout-smoke, using synthetic files."""
from pathlib import Path
import sys
from zipfile import ZipFile
from docx import Document
from openpyxl import load_workbook

root = Path(sys.argv[1]).resolve(strict=True)
word_path = root / '编辑后-段落测试.docx'
sheet_path = root / '编辑后-单元格测试.xlsx'
word = Document(word_path)
assert word.paragraphs[0].text == '重新输入标题'
assert word.paragraphs[1].text == '请检查端口配置与连接方式。并确认状态。'
assert any(r.text == '端口' and r.bold for r in word.paragraphs[1].runs)
assert word.tables[0].cell(0, 1).text == '表格内直接编辑'
assert word.sections[0].header.paragraphs[0].text == '网络运维资料'
assert len(word.inline_shapes) == 1
with ZipFile(root / '段落测试.docx') as before, ZipFile(word_path) as after:
    for name in before.namelist():
        if name != 'word/document.xml':
            assert before.read(name) == after.read(name), name
    assert b'bookmarkStart' not in after.read('word/document.xml')
sheet = load_workbook(sheet_path, data_only=False)
assert sheet['设备清单']['A2'].value == '网格中的新名称'
assert sheet['设备清单']['B2'].value == '新空白格'
assert sheet['设备清单']['B3'].value == '输入栏创建的值'
assert sheet['设备清单']['D4'].value == '重启后恢复的格子'
assert sheet['设备清单']['AA120'].value == '远端编辑已定位'
assert str(next(iter(sheet['设备清单'].merged_cells.ranges))) == 'A1:B1'
assert sheet['预算']['B2'].value == 150
assert sheet['预算']['D2'].value == '=SUM(B2:C2)'
assert sheet['预算'].sheet_state == 'hidden'
assert load_workbook(sheet_path, data_only=True)['预算']['D2'].value is None
print('Independent Word/Excel page/grid export verification passed.')
