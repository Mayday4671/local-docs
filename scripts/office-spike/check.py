from pathlib import Path
from docx import Document
from zipfile import ZipFile
from lxml import etree
import json, hashlib
root=Path('test-results/office-spike')
a=Document(root/'fixture.docx'); b=Document(root/'edited.docx')
checks={
 'edited_text': any('已完成离线修改' in p.text for p in b.paragraphs),
 'old_marker_removed': all('ORIGINAL_MARKER' not in p.text for p in b.paragraphs),
 'tables': [[[c.text for c in r.cells] for r in t.rows] for t in a.tables] == [[[c.text for c in r.cells] for r in t.rows] for t in b.tables],
 'header': [p.text for p in a.sections[0].header.paragraphs] == [p.text for p in b.sections[0].header.paragraphs],
 'footer': [p.text for p in a.sections[0].footer.paragraphs] == [p.text for p in b.sections[0].footer.paragraphs],
 'picture_count': len(a.inline_shapes)==len(b.inline_shapes)==1,
 'bold_text': any(r.text=='验收样本' and r.bold for p in b.paragraphs for r in p.runs),
 'unmodified_second_page': any(p.text=='未修改的第二页内容。' for p in b.paragraphs),
}
with ZipFile(root/'fixture.docx') as za, ZipFile(root/'edited.docx') as zb:
    images=lambda z: sorted(hashlib.sha256(z.read(p)).hexdigest() for p in z.namelist() if p.startswith('word/media/'))
    checks['image_bytes'] = images(za) == images(zb)
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    page_breaks=lambda z: len(etree.fromstring(z.read('word/document.xml')).xpath('//w:br[@w:type="page"]',namespaces=ns))
    checks['explicit_page_breaks'] = page_breaks(za)==page_breaks(zb)==1
    for name in ['word/styles.xml','word/header1.xml','word/footer1.xml']:
        checks[name+'_unchanged'] = za.read(name)==zb.read(name)
semantic=lambda paragraphs: [(p.text, p.style.style_id, p.alignment, [(r.text, r.bold, r.italic, r.underline, r.font.name, r.font.size) for r in p.runs]) for p in paragraphs]
checks['header_semantics']=semantic(a.sections[0].header.paragraphs)==semantic(b.sections[0].header.paragraphs)
checks['footer_semantics']=semantic(a.sections[0].footer.paragraphs)==semantic(b.sections[0].footer.paragraphs)
checks['page_geometry'] = [(s.page_width,s.page_height,s.top_margin,s.bottom_margin,s.left_margin,s.right_margin) for s in a.sections] == [(s.page_width,s.page_height,s.top_margin,s.bottom_margin,s.left_margin,s.right_margin) for s in b.sections]
(root/'roundtrip-checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
# Header/footer XML is rewritten: keep that evidence, and separately verify content and formatting.
assert all(v for k,v in checks.items() if k not in ['word/header1.xml_unchanged','word/footer1.xml_unchanged'])
