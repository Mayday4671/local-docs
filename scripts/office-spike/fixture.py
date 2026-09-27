from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw
from pathlib import Path

root = Path(__file__).parent
doc = Document()
doc.styles['Normal'].font.name = 'Microsoft YaHei'
doc.styles['Normal'].font.size = Pt(11)
doc.add_heading('离线编辑验证', 0)
paragraph = doc.add_paragraph('项目名称：')
paragraph.add_run('验收样本').bold = True
doc.add_paragraph('待修改内容 ORIGINAL_MARKER。')
table = doc.add_table(rows=3, cols=2)
table.style = 'Table Grid'
for row, values in zip(table.rows, [('项目', '数量'), ('本地备份', '12'), ('中文资料', '8')]):
    for cell, value in zip(row.cells, values): cell.text = value
table.rows[0].cells[0].paragraphs[0].runs[0].bold = True
doc.sections[0].header.paragraphs[0].text = '仅供离线测试的页眉'
doc.sections[0].footer.paragraphs[0].text = '保留页脚'
im = Image.new('RGB', (120, 48), '#087eff')
ImageDraw.Draw(im).rectangle((12, 12, 106, 34), fill='white')
im.save(root / 'fixture.png')
doc.add_picture(str(root / 'fixture.png'), width=Inches(1.2))
doc.add_page_break()
doc.add_heading('第二页', level=1)
doc.add_paragraph('未修改的第二页内容。')
doc.save(root / 'fixture.docx')
