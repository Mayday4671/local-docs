"""Read the synthetic desktop regression export with an independent DOCX implementation."""
from pathlib import Path
import sys
from zipfile import ZipFile
from docx import Document

root = Path(sys.argv[1]).resolve(strict=True)
doc = Document(root / '修复后.docx')
assert doc.paragraphs[0].text == '标题已可正常输入'
assert doc.tables[0].cell(1, 1).text == '保留的最终备注'
assert len(doc.inline_shapes) == 1
assert len(doc.paragraphs[0].paragraph_format.tab_stops) == 1
assert len(doc.tables[0].cell(1, 1).paragraphs[0].paragraph_format.tab_stops) == 1
with ZipFile(root / '段落测试.docx') as before, ZipFile(root / '修复后.docx') as after:
    for name in before.namelist():
        if name != 'word/document.xml':
            assert before.read(name) == after.read(name), name
print('Independent DOCX export check passed: title, formerly empty table cell, tab stops and untouched parts.')
