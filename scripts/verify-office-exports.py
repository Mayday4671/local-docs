"""Independent OOXML interoperability check for content-smoke exports (synthetic data)."""
from pathlib import Path
import sys
from docx import Document
from openpyxl import load_workbook

root = Path(sys.argv[1]).resolve(strict=True)
word = Document(root / '已编辑.docx')
assert word.paragraphs[0].text == '已编辑的项目标题'
assert word.tables[0].cell(0, 1).text == '交换机'
assert word.sections[0].header.paragraphs[0].text == '网络运维资料'
assert len(word.inline_shapes) == 1
sheet = load_workbook(root / '已编辑.xlsx', data_only=False)
assert sheet['设备清单']['A2'].value == '已编辑的交换机'
assert sheet['预算']['D2'].value == '=SUM(B2:C2)'
assert str(next(iter(sheet['设备清单'].merged_cells.ranges))) == 'A1:B1'
assert load_workbook(root / '已编辑.xlsx', data_only=True)['预算']['D2'].value is None
new_word = Document(root / '新建Word.docx')
assert [p.text for p in new_word.paragraphs] == ['新建文件的正文', '第二段正文', '第三段正文']
new_sheet = load_workbook(root / '新建Excel.xlsx')
assert new_sheet['工作表1']['A1'].value == '新建文件的正文'
print('Independent python-docx/openpyxl export checks passed.')
